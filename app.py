"""
Pro Resignation Approval System  (Streamlit)
Run:      streamlit run resignation_pro.py
Install:  pip install streamlit pandas python-docx xlsxwriter

Optional .streamlit/secrets.toml:
    APP_PASSWORD = "choose-a-password"     # enables login screen
    SMTP_HOST = "smtp.gmail.com"
    SMTP_USER = "you@company.com"
    SMTP_PASS = "app-password"
"""
import base64, hmac, html, io, json, os, smtplib, sqlite3
from datetime import date, datetime, timedelta
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from docx import Document

DB_PATH = os.getenv("RESIGN_DB", "resignations.db")
STATUSES = ["Pending", "Approved", "Rejected"]
STATUS_COLOR = {"Pending": "#C98A1B", "Approved": "#2E7D32", "Rejected": "#B3261E"}
NOTICE = {"Waived / None": 0, "2 Weeks": 14, "One Month": 30, "Two Months": 60}
CHECKLIST = ["Laptop / equipment returned", "ID badge returned", "Keys returned",
             "System access revoked", "Knowledge handover done",
             "Exit interview done", "Final dues approved by HR"]
EARLY_LEAVER_DAYS = 90

st.set_page_config(page_title="Resignation Approval System", page_icon="📄", layout="wide")


# ----------------------------------------------------------------- helpers
def secret(name, default=""):
    try:
        return st.secrets.get(name, os.getenv(name, default))
    except Exception:
        return os.getenv(name, default)


def to_english_str(text):
    return str(text or "").translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")).strip()


def esc(v):
    return html.escape(str(v if v is not None else ""))


def D(s):
    return date.fromisoformat(s) if isinstance(s, str) else s


def fmt(s, f="%d/%m/%Y"):
    return D(s).strftime(f)


# ---------------------------------------------------------------- database
def run(sql, args=(), fetch=None):
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    try:
        cur = con.execute(sql, args)
        con.commit()
        if fetch == "all":
            return [dict(r) for r in cur.fetchall()]
        if fetch == "one":
            r = cur.fetchone()
            return dict(r) if r else None
        return cur.lastrowid
    finally:
        con.close()


def init_db():
    run("""CREATE TABLE IF NOT EXISTS resignations(
        id INTEGER PRIMARY KEY AUTOINCREMENT, recipient_name TEXT, employee_name TEXT,
        employee_code TEXT, job_title TEXT, branch TEXT, grading TEXT,
        hiring_date TEXT, resignation_date TEXT, last_working_day TEXT,
        final_dues REAL DEFAULT 0, loan_balance REAL DEFAULT 0, other_deductions REAL DEFAULT 0,
        reason TEXT, direct_manager TEXT, notice_period TEXT, handover_status TEXT,
        absence_days INTEGER DEFAULT 0, status TEXT DEFAULT 'Pending', manager_comment TEXT,
        checklist TEXT DEFAULT '[]', photo BLOB, created_at TEXT, updated_at TEXT)""")
    run("""CREATE TABLE IF NOT EXISTS audit(
        id INTEGER PRIMARY KEY AUTOINCREMENT, rid INTEGER, action TEXT, detail TEXT, ts TEXT)""")
    if not run("SELECT COUNT(*) n FROM resignations", fetch="one")["n"]:
        save(dict(recipient_name="Mohamed", employee_name="Abu Bakr Muhammad Mahmoud",
                  employee_code="E-1001", job_title="Sales Casual", branch="Park-St Branch",
                  grading="G1 - L1", hiring_date="2026-10-03", resignation_date="2026-10-20",
                  last_working_day="2026-10-20", final_dues=7000, loan_balance=0,
                  other_deductions=0, reason="Education", direct_manager="Sherif S. El-Sayed",
                  notice_period="One Month", handover_status="Team Leader", absence_days=0,
                  status="Pending", manager_comment="", checklist="[]"))


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def log(rid, action, detail=""):
    run("INSERT INTO audit(rid,action,detail,ts) VALUES(?,?,?,?)", (rid, action, detail, now()))


def get_all():
    return run("SELECT * FROM resignations ORDER BY id DESC", fetch="all")


def get_one(rid):
    return run("SELECT * FROM resignations WHERE id=?", (rid,), fetch="one")


def save(rec, rid=None):
    rec = dict(rec, updated_at=now())
    if rid:
        run(f"UPDATE resignations SET {', '.join(c + '=?' for c in rec)} WHERE id=?", [*rec.values(), rid])
        log(rid, "Edited", rec.get("employee_name", ""))
        return rid
    rec["created_at"] = now()
    new = run(f"INSERT INTO resignations({','.join(rec)}) VALUES({','.join('?' * len(rec))})", list(rec.values()))
    log(new, "Created", rec.get("employee_name", ""))
    return new


# ------------------------------------------------------------ calculations
def calc(d):
    hire, res, last = D(d["hiring_date"]), D(d["resignation_date"]), D(d["last_working_day"])
    notice_days = NOTICE.get(d["notice_period"], 0)
    required_last = res + timedelta(days=notice_days)
    recipient = (d.get("recipient_name") or "").strip() or "Manager"
    if not recipient.lower().startswith(("mr.", "ms.", "mrs.")):
        recipient = f"Mr. {recipient}"
    return dict(
        exp=max((last - hire).days, 0),
        net=(d["final_dues"] or 0) - (d["loan_balance"] or 0) - (d["other_deductions"] or 0),
        required_last=required_last,
        shortfall=max((required_last - last).days, 0),
        recipient=recipient,
        done=json.loads(d.get("checklist") or "[]"),
    )


def rows_for(d, c):
    return [("Employee Name", d["employee_name"]), ("Employee Code", d["employee_code"]),
            ("Branch", d["branch"]), ("Job Title", d["job_title"]), ("Grading", d["grading"]),
            ("Hiring Date", fmt(d["hiring_date"])), ("Resignation Date", fmt(d["resignation_date"])),
            ("Last Working Day", fmt(d["last_working_day"])), ("Total Experience", f"{c['exp']} Days"),
            ("Notice Period", d["notice_period"]),
            ("Notice Shortfall", f"{c['shortfall']} Days" if c["shortfall"] else "None"),
            ("Final Dues (EGP)", f"{d['final_dues']:,.0f}"), ("Loan Deduction (EGP)", f"{d['loan_balance']:,.0f}"),
            ("Other Deductions (EGP)", f"{d['other_deductions']:,.0f}"), ("Total Net Pay (EGP)", f"{c['net']:,.0f}"),
            ("Reason for Leaving", d["reason"]), ("Direct Manager", d["direct_manager"]),
            ("Handover To", d["handover_status"]), ("Absence During Year", f"{d['absence_days']} Days"),
            ("Status", d["status"]), ("Manager Comment", d["manager_comment"] or "-")]


# ----------------------------------------------------------------- exports
def make_docx(d, c, company):
    doc = Document()
    doc.add_heading(f"{company} - Resignation Approval Request", 0)
    doc.add_paragraph(f"Dear {c['recipient']},")
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in rows_for(d, c):
        cells = t.add_row().cells
        cells[0].text, cells[1].text = k, str(v)
    doc.add_heading("Handover checklist", 2)
    for item in CHECKLIST:
        doc.add_paragraph(("[x] " if item in c["done"] else "[  ] ") + item)
    doc.add_paragraph("\n\nEmployee signature: ______________    Manager signature: ______________    HR: ______________")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def make_excel(df):
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="xlsxwriter") as w:
        df.to_excel(w, index=False, sheet_name="Resignations")
        ws = w.sheets["Resignations"]
        for i, col in enumerate(df.columns):
            ws.set_column(i, i, min(max(len(str(col)), *(len(str(x)) for x in df[col])) + 2, 40) if len(df) else 15)
        ws.freeze_panes(1, 0)
    return bio.getvalue()


def to_df(rows):
    out = []
    for d in rows:
        c = calc(d)
        out.append({"ID": d["id"], "Employee": d["employee_name"], "Code": d["employee_code"],
                    "Branch": d["branch"], "Job Title": d["job_title"], "Hired": d["hiring_date"],
                    "Resigned": d["resignation_date"], "Last Day": d["last_working_day"],
                    "Experience (days)": c["exp"], "Net Pay (EGP)": c["net"], "Reason": d["reason"],
                    "Manager": d["direct_manager"], "Notice": d["notice_period"],
                    "Notice Shortfall (days)": c["shortfall"], "Status": d["status"],
                    "Checklist": f"{len(c['done'])}/{len(CHECKLIST)}"})
    return pd.DataFrame(out)


# --------------------------------------------------------------- THEME + A4 card UI
PALETTES = {
    "Espresso": dict(main="#8C5A2B", dark="#5E3A1A", soft="#F3E9DD", bg="#FAF7F2", line="#E3D3BF"),
    "Ocean": dict(main="#1F6FEB", dark="#134B9E", soft="#E6F0FF", bg="#F4F8FD", line="#CFE0FA"),
    "Emerald": dict(main="#12805C", dark="#0B5A41", soft="#E3F5EE", bg="#F3F9F6", line="#C9E8DC"),
    "Royal": dict(main="#6D3FD1", dark="#4A2699", soft="#EEE7FB", bg="#F7F4FD", line="#DCCFF6"),
    "Rose": dict(main="#D1325B", dark="#9B1F42", soft="#FCE6EC", bg="#FCF5F7", line="#F5CFD9"),
}
P = PALETTES["Espresso"]


def _tok(css, p):
    for k, t in (("main", "__M__"), ("dark", "__D__"), ("soft", "__S__"), ("bg", "__BG__"), ("line", "__L__")):
        css = css.replace(t, p[k])
    return css


THEME_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&display=swap');
html,body,[class*="css"],.stApp{font-family:'Inter',sans-serif}
.stApp{background:radial-gradient(1100px 480px at 92% -8%,__S__ 0%,transparent 62%),__BG__}
#MainMenu,footer{visibility:hidden}
header[data-testid="stHeader"]{background:transparent}
.block-container{padding-top:1.6rem;max-width:1250px;animation:up .5s ease}
@keyframes up{from{opacity:0;transform:translateY(10px)}to{opacity:1;transform:none}}
h1,h2,h3,h4{color:__D__;font-weight:800}

.hero{display:flex;justify-content:space-between;align-items:center;gap:24px;flex-wrap:wrap;color:#fff;
 background:linear-gradient(120deg,__D__,__M__ 72%);border-radius:22px;padding:26px 32px;margin-bottom:20px;
 box-shadow:0 18px 38px -16px __M__;position:relative;overflow:hidden}
.hero:after{content:"";position:absolute;right:-70px;top:-90px;width:280px;height:280px;border-radius:50%;background:rgba(255,255,255,.09)}
.hero:before{content:"";position:absolute;right:140px;bottom:-110px;width:200px;height:200px;border-radius:50%;background:rgba(255,255,255,.06)}
.hero-sub{font-size:12px;letter-spacing:.2em;opacity:.8;font-weight:700}
.hero-title{font-size:30px;font-weight:800;line-height:1.15;margin:4px 0}
.hero-note{font-size:13px;opacity:.85}
.chips{display:flex;gap:10px;position:relative;z-index:1}
.chip{background:rgba(255,255,255,.16);border:1px solid rgba(255,255,255,.28);border-radius:16px;padding:10px 18px;text-align:center;min-width:82px}
.chip b{display:block;font-size:26px;font-weight:800}
.chip span{font-size:11px;opacity:.9;text-transform:uppercase;letter-spacing:.08em}

.stTabs [data-baseweb="tab-list"]{gap:6px;background:#fff;padding:6px;border-radius:16px;border:1px solid __L__;
 box-shadow:0 8px 20px -12px rgba(0,0,0,.3);width:fit-content}
.stTabs [data-baseweb="tab-highlight"],.stTabs [data-baseweb="tab-border"]{display:none}
.stTabs [data-baseweb="tab"]{height:42px;border-radius:12px;padding:0 22px;font-weight:700;color:__M__;background:transparent;transition:all .2s}
.stTabs [data-baseweb="tab"]:hover{background:__S__}
.stTabs [aria-selected="true"]{background:linear-gradient(135deg,__M__,__D__)!important;box-shadow:0 8px 18px -8px __M__}
.stTabs [aria-selected="true"],.stTabs [aria-selected="true"] *{color:#fff!important}
.stTabs [data-baseweb="tab-panel"]{padding-top:22px}

[data-testid="stMetric"]{background:#fff;border:1px solid __L__;border-left:5px solid __M__;border-radius:14px;padding:14px 16px;
 box-shadow:0 10px 22px -16px rgba(0,0,0,.35);transition:transform .2s}
[data-testid="stMetric"]:hover{transform:translateY(-3px)}
[data-testid="stMetricValue"]{color:__D__;font-weight:800}
[data-testid="stExpander"],[data-testid="stForm"]{background:#fff;border:1px solid __L__!important;border-radius:16px;
 box-shadow:0 10px 24px -18px rgba(0,0,0,.35)}
[data-testid="stDataFrame"]{border:1px solid __L__;border-radius:14px;overflow:hidden}
section[data-testid="stSidebar"]{background:#fff;border-right:1px solid __L__}

.stButton>button,.stDownloadButton>button,.stFormSubmitButton>button{border-radius:12px;font-weight:700;border:1.5px solid __L__;transition:all .2s}
.stButton>button:hover,.stDownloadButton>button:hover{border-color:__M__;color:__M__;transform:translateY(-1px)}
button[kind="primary"],button[kind="primaryFormSubmit"],[data-testid="stBaseButton-primary"],[data-testid="stBaseButton-primaryFormSubmit"]{
 background:linear-gradient(135deg,__M__,__D__)!important;color:#fff!important;border:0!important;box-shadow:0 8px 18px -8px __M__}
[data-baseweb="input"],[data-baseweb="select"]>div,[data-baseweb="textarea"]{border-radius:10px!important}
[data-baseweb="input"]:focus-within,[data-baseweb="select"]>div:focus-within{border-color:__M__!important;box-shadow:0 0 0 3px __S__}

.pillrow{display:flex;align-items:center;gap:12px;margin:2px 0 14px}
.pill{color:#fff;border-radius:999px;padding:4px 14px;font-size:12px;font-weight:700;letter-spacing:.05em}
.meta{color:#777;font-size:12px}
"""

CARD_CSS = """
body{font-family:'Inter',sans-serif;background:transparent;margin:0;padding:8px;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.btns{display:flex;gap:10px;justify-content:flex-end;margin-bottom:12px}
.btn{background:linear-gradient(135deg,__M__,__D__);color:#fff;border:0;padding:10px 18px;border-radius:10px;font-weight:700;cursor:pointer;box-shadow:0 6px 14px -6px __M__}
.card{background:#fff;border:1px solid __L__;border-radius:18px;max-width:800px;margin:0 auto;overflow:hidden;box-shadow:0 18px 40px -22px rgba(0,0,0,.35)}
.band{background:linear-gradient(120deg,__D__,__M__);color:#fff;padding:22px 26px;display:flex;justify-content:space-between;gap:14px;align-items:center}
.band small{letter-spacing:.16em;font-size:10px;opacity:.8;font-weight:700}
.band h1{margin:4px 0 10px;font-size:24px;line-height:1.12;font-weight:800}
.pillc{display:inline-block;background:#fff;border-radius:999px;padding:3px 12px;font-size:11px;font-weight:800}
.prof{display:flex;gap:12px;align-items:center;background:rgba(255,255,255,.14);border:1px solid rgba(255,255,255,.3);border-radius:14px;padding:10px}
.body{padding:20px 26px 24px}
.tl{display:flex;position:relative;margin:4px 0 18px}
.tl:before{content:"";position:absolute;left:10%;right:10%;top:15px;height:3px;background:__L__}
.step{flex:1;text-align:center;position:relative;font-size:11px}
.dot{width:30px;height:30px;border-radius:50%;margin:0 auto 5px;background:__M__;color:#fff;display:flex;align-items:center;justify-content:center;font-size:14px;border:3px solid #fff;box-shadow:0 0 0 2px __M__}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:10px}
.kpi{background:__S__;border-radius:12px;padding:10px 12px}
.kpi small{display:block;font-size:10px;text-transform:uppercase;letter-spacing:.08em;color:__D__;font-weight:700;margin-bottom:2px}
.kpi b{font-size:15px;color:#111}
table{width:100%;border-collapse:collapse;margin-top:14px;border-radius:10px;overflow:hidden}
th{background:__M__;color:#fff;padding:7px 12px;text-align:left;font-size:12px}
td{padding:6px 12px;border-bottom:1px solid __L__;font-size:12px}
tr:nth-child(even) td{background:__S__}
.chk{margin-top:14px;border:1px solid __L__;border-radius:12px;padding:12px 14px}
.prog{height:8px;background:__S__;border-radius:99px;overflow:hidden;margin:6px 0 10px}
.prog i{display:block;height:100%;background:linear-gradient(90deg,__M__,__D__)}
.sig{display:flex;gap:24px;margin-top:30px;font-size:11px;color:#666}.sig div{flex:1;border-top:1px solid #aaa;padding-top:4px}
@media print{.no-print{display:none!important}body{padding:0}.card{box-shadow:none;border:0;max-width:100%}}
"""


def theme_css(p):
    return _tok(THEME_CSS, p)


def card_css():
    return _tok(CARD_CSS, P)


def card_html(d, c, company):
    color = STATUS_COLOR[d["status"]]
    if d.get("photo"):
        mime = "image/png" if d["photo"][:4] == b"\x89PNG" else "image/jpeg"
        img = (f'<img src="data:{mime};base64,{base64.b64encode(d["photo"]).decode()}" '
               'style="width:84px;height:110px;object-fit:cover;border-radius:10px;border:2px solid #fff">')
    else:
        img = ('<div style="width:84px;height:110px;background:rgba(255,255,255,.2);border-radius:10px;'
               'display:flex;align-items:center;justify-content:center;font-size:36px">👤</div>')
    steps = [("👤", "Joined", fmt(d["hiring_date"])), ("🏆", "Experience", f"{c['exp']} Days"),
             ("📄", "Submitted", fmt(d["resignation_date"])), ("🕒", "Notice", d["notice_period"]),
             ("✅", "Status", d["status"])]
    tl = "".join(f'<div class="step"><div class="dot">{i}</div><b>{esc(t)}</b><div style="color:#555">{esc(v)}</div></div>'
                 for i, t, v in steps)
    kpis = [("Net Pay (EGP)", f"{c['net']:,.0f}"), ("Last Day", fmt(d["last_working_day"])),
            ("Reason", d["reason"]), ("Manager", d["direct_manager"])]
    kpi_html = "".join(f'<div class="kpi"><small>{esc(t)}</small><b>{esc(v)}</b></div>' for t, v in kpis)
    table = "".join(f"<tr><td style='width:38%'>{esc(k)}</td><td><b>{esc(v)}</b></td></tr>" for k, v in rows_for(d, c))
    todo = "".join(f"<div>{'☑' if i in c['done'] else '☐'} {esc(i)}</div>" for i in CHECKLIST)
    pct = int(100 * len(c["done"]) / len(CHECKLIST))
    warn = (f'<div style="margin-top:10px;color:#B3261E;font-weight:700;font-size:12px">⚠ Notice shortfall: '
            f'{c["shortfall"]} days (required last day {fmt(c["required_last"])})</div>') if c["shortfall"] else ""
    js = """<script>function dl(){html2canvas(document.getElementById('cap'),{scale:2}).then(c=>{const a=document.createElement('a');
    a.download='Resignation___NAME__.png';a.href=c.toDataURL('image/png');a.click();});}</script>""".replace(
        "__NAME__", esc(d["employee_name"]).replace(" ", "_"))
    return f"""<html><head><script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script>
<style>{card_css()}</style></head><body>
<div class="btns no-print"><button class="btn" onclick="dl()">🖼️ PNG</button><button class="btn" onclick="window.print()">🖨️ Print / PDF</button></div>
<div id="cap" class="card">
<div class="band"><div><small>{esc(company).upper()}</small><h1>Resignation<br>Approval Request</h1>
<span class="pillc" style="color:{color}">● {esc(d['status']).upper()}</span>
<div style="font-size:12px;opacity:.9;margin-top:10px">Dear {esc(c['recipient'])},</div></div>
<div class="prof">{img}<div><small>EMPLOYEE</small><div style="font-size:17px;font-weight:800">{esc(d['employee_name'])}</div>
<div style="font-size:12px;opacity:.9">{esc(d['job_title'])} · {esc(d['employee_code'])}</div>
<div style="font-size:11px;opacity:.8;margin-top:4px">{esc(d['branch'])}</div></div></div></div>
<div class="body"><div class="tl">{tl}</div><div class="kpis">{kpi_html}</div>{warn}
<table><thead><tr><th>Field</th><th>Value</th></tr></thead><tbody>{table}</tbody></table>
<div class="chk"><div style="color:{P['main']};font-weight:800;font-size:13px">HANDOVER CHECKLIST · {len(c['done'])}/{len(CHECKLIST)}</div>
<div class="prog"><i style="width:{pct}%"></i></div>
<div style="display:grid;grid-template-columns:1fr 1fr;gap:4px;font-size:11px">{todo}</div></div>
<div class="sig"><div>Employee</div><div>Direct Manager</div><div>HR</div></div>
</div></div>{js}</body></html>"""


# --------------------------------------------------------------------- app
def gate():
    pw = secret("APP_PASSWORD")
    if not pw or st.session_state.get("auth"):
        return
    st.title("🔒 Resignation Approval System")
    entered = st.text_input("Password", type="password")
    if st.button("Login", type="primary"):
        if hmac.compare_digest(entered, pw):
            st.session_state.auth = True
            st.rerun()
        st.error("Wrong password")
    st.stop()


init_db()
P = PALETTES[st.session_state.get("theme", "Espresso")]
st.markdown(f"<style>{theme_css(P)}</style>", unsafe_allow_html=True)
gate()
if "flash" in st.session_state:
    st.toast(st.session_state.pop("flash"))

with st.sidebar:
    st.header("⚙️ Settings")
    company = st.text_input("Company name", value="Company")
    st.selectbox("🎨 Theme color", list(PALETTES), key="theme")
    st.caption("Data is stored in " + DB_PATH)
    if os.path.exists(DB_PATH):
        st.download_button("💾 Backup database", open(DB_PATH, "rb").read(), "resignations_backup.db")

_all = get_all()
_cnt = [("Total", len(_all))] + [(k, sum(1 for r in _all if r["status"] == k)) for k in STATUSES]
st.markdown(
    f'<div class="hero"><div><div class="hero-sub">{esc(company).upper()}</div>'
    '<div class="hero-title">Resignation Approval System</div>'
    '<div class="hero-note">Track, approve and archive every resignation in one place</div></div>'
    '<div class="chips">' + "".join(f'<div class="chip"><b>{v}</b><span>{k}</span></div>' for k, v in _cnt)
    + '</div></div>', unsafe_allow_html=True)

tab_view, tab_manage, tab_dash, tab_audit = st.tabs(
    ["📄  Requests", "✏️  Manage", "📊  Analytics", "🕘  Audit Log"])

# ------------------------------------------------------------- TAB 1: view
with tab_view:
    rows = get_all()
    if not rows:
        st.warning("⚠️ لا توجد بيانات استقالات مسجلة حالياً.")
    else:
        f1, f2 = st.columns([1, 2])
        status_filter = f1.multiselect("Filter by status", STATUSES, default=STATUSES)
        pool = [r for r in rows if r["status"] in status_filter]
        if not pool:
            st.info("No records match this filter.")
        else:
            labels = {r["id"]: f"{r['employee_name']} — {r['branch']} [{r['status']}]" for r in pool}
            sel = f2.selectbox("اختر الموظف:", list(labels), format_func=labels.get)
            d = get_one(sel)
            c = calc(d)
            st.markdown(
                f'<div class="pillrow"><span class="pill" style="background:{STATUS_COLOR[d["status"]]}">{d["status"].upper()}</span>'
                f'<span class="meta">Last updated {esc(d["updated_at"])}</span></div>', unsafe_allow_html=True)

            with st.expander("✅ Decision & Handover Checklist", expanded=True):
                k1, k2 = st.columns([1, 2])
                new_status = k1.selectbox("Decision", STATUSES, index=STATUSES.index(d["status"]), key=f"st_{sel}")
                comment = k2.text_input("Manager comment", value=d["manager_comment"] or "", key=f"cm_{sel}")
                done = st.multiselect("Completed items", CHECKLIST, default=c["done"], key=f"ck_{sel}")
                if st.button("💾 Save decision", type="primary", key=f"sv_{sel}"):
                    run("UPDATE resignations SET status=?, manager_comment=?, checklist=?, updated_at=? WHERE id=?",
                        (new_status, comment, json.dumps(done), now(), sel))
                    log(sel, "Decision", f"{d['status']} → {new_status}; checklist {len(done)}/{len(CHECKLIST)}")
                    st.session_state.flash = "Decision saved ✅"
                    st.rerun()

            if c["shortfall"]:
                st.warning(f"Notice shortfall of {c['shortfall']} days — required last working day is {fmt(c['required_last'])}.")

            st.markdown("#### 📥 Export")
            e1, e2, e3 = st.columns(3)
            safe = d["employee_name"].replace(" ", "_")
            word = make_docx(d, c, company)
            e1.download_button("📄 Word (.docx)", word, f"Resignation_{safe}.docx", use_container_width=True,
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            e2.download_button("📊 Excel (this employee)", make_excel(to_df([d])), f"Resignation_{safe}.xlsx",
                               use_container_width=True,
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            e3.download_button("📚 Excel (all records)", make_excel(to_df(rows)), "All_Resignations.xlsx",
                               use_container_width=True,
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

            components.html(card_html(d, c, company), height=1120, scrolling=True)

            with st.expander("📧 Send by email (Word file attached)"):
                a, b = st.columns(2)
                sender = a.text_input("Sender", value=secret("SMTP_USER"))
                pwd = a.text_input("App password", type="password", value=secret("SMTP_PASS"))
                receiver = b.text_input("Receiver")
                host = b.text_input("SMTP host", value=secret("SMTP_HOST", "smtp.gmail.com"))
                if st.button("🚀 Send"):
                    if not (sender and pwd and receiver):
                        st.error("Please fill in sender, password and receiver.")
                    else:
                        try:
                            msg = MIMEMultipart()
                            msg["Subject"] = f"Resignation Approval Request - {d['employee_name']}"
                            msg["From"], msg["To"] = sender, receiver
                            lines = "".join(f"<p><b>{esc(k)}:</b> {esc(v)}</p>" for k, v in rows_for(d, c))
                            msg.attach(MIMEText(f"<h3>Dear {esc(c['recipient'])},</h3>{lines}", "html"))
                            part = MIMEApplication(word, Name=f"Resignation_{safe}.docx")
                            part["Content-Disposition"] = f'attachment; filename="Resignation_{safe}.docx"'
                            msg.attach(part)
                            with smtplib.SMTP(host, 587, timeout=30) as s:
                                s.starttls()
                                s.login(sender, pwd)
                                s.sendmail(sender, receiver, msg.as_string())
                            log(sel, "Emailed", receiver)
                            st.success("✅ Email sent!")
                        except Exception as ex:
                            st.error(f"❌ Failed: {ex}")

# ------------------------------------------------------------ TAB 2: manage
with tab_manage:
    st.subheader("⚙️ إدارة تسجيلات الاستقالات")
    action = st.radio("الإجراء", ["إضافة ➕", "تعديل ✏️", "حذف 🗑️"], horizontal=True)
    rows = get_all()
    labels = {r["id"]: f"{r['employee_name']} — {r['branch']}" for r in rows}

    if action == "حذف 🗑️":
        if not rows:
            st.warning("لا توجد استقالات.")
        else:
            del_id = st.selectbox("اختر السجل", list(labels), format_func=labels.get)
            sure = st.checkbox("I understand this cannot be undone")
            if st.button("❌ Delete permanently", type="primary", disabled=not sure):
                run("DELETE FROM resignations WHERE id=?", (del_id,))
                log(del_id, "Deleted", labels[del_id])
                st.session_state.flash = "Deleted"
                st.rerun()
    else:
        edit = action == "تعديل ✏️"
        rid = None
        t = dict(recipient_name="", employee_name="", employee_code="", job_title="", branch="", grading="",
                 hiring_date=date.today(), resignation_date=date.today(), last_working_day=date.today(),
                 final_dues=0.0, loan_balance=0.0, other_deductions=0.0, reason="", direct_manager="",
                 notice_period="One Month", handover_status="", absence_days=0, status="Pending",
                 manager_comment="", checklist="[]", photo=None)
        if edit:
            if not rows:
                st.warning("لا توجد سجلات للتعديل.")
                st.stop()
            rid = st.selectbox("اختر الموظف", list(labels), format_func=labels.get)
            t = get_one(rid)
        key = lambda n: f"{n}_{rid or 'new'}"
        with st.form(key("form")):
            r1, r2 = st.columns(2)
            recipient = r1.text_input("المرسل إليه (Recipient)", t["recipient_name"], key=key("rec"))
            manager = r2.text_input("المدير المباشر (Direct Manager)", t["direct_manager"], key=key("mgr"))
            c1, c2 = st.columns(2)
            name = c1.text_input("اسم الموظف *", t["employee_name"], key=key("nm"))
            code = c1.text_input("كود الموظف (Employee Code)", t["employee_code"], key=key("cd"))
            job = c1.text_input("المسمى الوظيفي", t["job_title"], key=key("jb"))
            branch = c1.text_input("الفرع", t["branch"], key=key("br"))
            grading = c1.text_input("الدرجة الوظيفية", t["grading"], key=key("gr"))
            reason = c1.text_input("سبب الاستقالة", t["reason"], key=key("rs"))
            hire = c2.date_input("تاريخ التعيين", D(t["hiring_date"]), key=key("hd"))
            resd = c2.date_input("تاريخ تقديم الاستقالة", D(t["resignation_date"]), key=key("rd"))
            last = c2.date_input("آخر يوم عمل", D(t["last_working_day"]), key=key("ld"))
            notice = c2.selectbox("فترة الإشعار", list(NOTICE),
                                  index=list(NOTICE).index(t["notice_period"]) if t["notice_period"] in NOTICE else 2,
                                  key=key("np"))
            handover = c2.text_input("التسليم إلى (Handover To)", t["handover_status"], key=key("ho"))
            c3, c4, c5, c6 = st.columns(4)
            dues = c3.number_input("المستحقات (EGP)", value=float(t["final_dues"]), step=500.0, key=key("du"))
            loan = c4.number_input("خصم السلفة", value=float(t["loan_balance"]), step=100.0, key=key("ln"))
            other = c5.number_input("خصومات أخرى", value=float(t["other_deductions"]), step=100.0, key=key("ot"))
            absence = c6.number_input("أيام الغياب", value=int(t["absence_days"]), step=1, key=key("ab"))
            photo = st.file_uploader("🖼️ صورة الموظف", type=["png", "jpg", "jpeg"], key=key("ph"))
            if st.form_submit_button("💾 حفظ", type="primary"):
                name_clean = to_english_str(name)
                dup = run("SELECT id FROM resignations WHERE employee_name=? AND hiring_date=? AND id IS NOT ?",
                          (name_clean, hire.isoformat(), rid), fetch="one")
                if not name_clean:
                    st.error("Employee name is required.")
                elif resd < hire:
                    st.error("Resignation date can't be before the hiring date.")
                elif last < resd:
                    st.error("Last working day can't be before the resignation date.")
                elif dup:
                    st.error("A record for this employee and hiring date already exists.")
                else:
                    rec = dict(recipient_name=to_english_str(recipient), employee_name=name_clean,
                               employee_code=to_english_str(code), job_title=to_english_str(job),
                               branch=to_english_str(branch), grading=to_english_str(grading),
                               hiring_date=hire.isoformat(), resignation_date=resd.isoformat(),
                               last_working_day=last.isoformat(), final_dues=dues, loan_balance=loan,
                               other_deductions=other, reason=to_english_str(reason),
                               direct_manager=to_english_str(manager), notice_period=notice,
                               handover_status=to_english_str(handover), absence_days=int(absence))
                    if photo is not None:
                        rec["photo"] = photo.getvalue()
                    save(rec, rid)
                    st.session_state.flash = "Saved ✅"
                    st.rerun()

# ----------------------------------------------------------- TAB 3: dashboard
with tab_dash:
    st.title("📊 داشبورد تحليل الاستقالات")
    rows = get_all()
    if not rows:
        st.info("لا توجد استقالات مسجلة.")
    else:
        df = to_df(rows)
        df["Last Day"] = pd.to_datetime(df["Last Day"])
        f1, f2, f3 = st.columns(3)
        sts = f1.multiselect("Status", STATUSES, default=STATUSES, key="d_st")
        brs = f2.multiselect("Branch", sorted(df["Branch"].unique()), key="d_br")
        lo, hi = df["Last Day"].min().date(), df["Last Day"].max().date()
        rng = f3.date_input("Last-day range", (lo, hi), key="d_rg")
        df = df[df["Status"].isin(sts)]
        if brs:
            df = df[df["Branch"].isin(brs)]
        if isinstance(rng, tuple) and len(rng) == 2:
            df = df[(df["Last Day"].dt.date >= rng[0]) & (df["Last Day"].dt.date <= rng[1])]

        if df.empty:
            st.info("No data for the selected filters.")
        else:
            early = (df["Experience (days)"] < EARLY_LEAVER_DAYS).mean() * 100
            m = st.columns(5)
            m[0].metric("Resignations", len(df))
            m[1].metric("Pending", int((df["Status"] == "Pending").sum()))
            m[2].metric("Total net pay (EGP)", f"{df['Net Pay (EGP)'].sum():,.0f}")
            m[3].metric("Avg tenure (days)", int(df["Experience (days)"].mean()))
            m[4].metric(f"Early leavers (<{EARLY_LEAVER_DAYS}d)", f"{early:.0f}%")
            st.markdown("---")
            a, b = st.columns(2)
            a.subheader("🏢 By branch")
            a.bar_chart(df["Branch"].value_counts(), color=P["main"])
            b.subheader("💡 Top reasons")
            b.bar_chart(df["Reason"].value_counts(), color=P["main"])
            c_, d_ = st.columns(2)
            c_.subheader("📈 Monthly trend")
            c_.line_chart(df.groupby(df["Last Day"].dt.to_period("M").astype(str)).size(), color=P["dark"])
            d_.subheader("👔 By manager")
            d_.bar_chart(df["Manager"].value_counts(), color=P["dark"])
            st.markdown("---")
            q = st.text_input("🔍 Search employee / branch / code").strip()
            view = df[df[["Employee", "Branch", "Code"]].apply(
                lambda col: col.astype(str).str.contains(q, case=False, na=False)).any(axis=1)] if q else df
            st.dataframe(view, use_container_width=True, hide_index=True)
            x1, x2 = st.columns(2)
            x1.download_button("📥 CSV", view.to_csv(index=False).encode("utf-8-sig"), "resignations_log.csv",
                               "text/csv", use_container_width=True)
            x2.download_button("📥 Excel", make_excel(view), "resignations_log.xlsx", use_container_width=True,
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ------------------------------------------------------------- TAB 4: audit
with tab_audit:
    st.subheader("🧾 Activity history")
    logs = run("SELECT ts AS Time, rid AS ID, action AS Action, detail AS Detail FROM audit ORDER BY id DESC LIMIT 500",
               fetch="all")
    if logs:
        st.dataframe(pd.DataFrame(logs), use_container_width=True, hide_index=True)
    else:
        st.info("No activity yet.")
